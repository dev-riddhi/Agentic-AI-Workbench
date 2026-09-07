"""Tool: Read / Write CSV, Excel, JSON, XML
Unified handler for reading and writing structured data in CSV, Excel (.xlsx), JSON, and XML formats.
"""

import csv
import io
import json
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET


from app.tools.file_security import resolve_safe_path


def read_write_csv_excel_json_xml(
    action: str,
    file_path: str,
    format: str | None = None,
    data: Any | None = None,
    sheet_name: str | None = None,
    root_tag: str = "root",
    encoding: str = "utf-8",
) -> dict[str, Any]:
    """Reads or writes structured datasets across CSV, Excel, JSON, and XML inside uploads.

    Args:
        action: 'read' or 'write'.
        file_path: Target file path inside the uploads directory.
        format: Format override ('csv', 'excel', 'json', 'xml'). If None, infers from extension.
        data: Data to write (list of dicts, dict, or nested structure).
        sheet_name: Specific sheet name for Excel operations.
        root_tag: XML root element tag name for XML writes (default: 'root').
        encoding: File character encoding (default: 'utf-8').

    Returns:
        dict: Result containing data or write status.
    """
    try:
        path = resolve_safe_path(file_path)
    except (PermissionError, ValueError) as err:
        return {
            "success": False,
            "error": str(err),
            "file_path": str(file_path),
        }

    act = action.strip().lower()

    # Determine format
    fmt = (format or "").lower().strip()
    if not fmt:
        ext = path.suffix.lower()
        if ext in (".csv",):
            fmt = "csv"
        elif ext in (".xlsx", ".xls"):
            fmt = "excel"
        elif ext in (".json",):
            fmt = "json"
        elif ext in (".xml",):
            fmt = "xml"
        else:
            return {
                "success": False,
                "error": f"Cannot infer data format from extension '{ext}'. Please specify 'format'.",
                "file_path": str(path),
            }

    if act == "read":
        return _read_data(path, fmt, sheet_name, encoding)
    elif act == "write":
        if data is None:
            return {
                "success": False,
                "error": "Parameter 'data' is required for write action.",
                "file_path": str(path),
            }
        return _write_data(path, fmt, data, sheet_name, root_tag, encoding)
    else:
        return {
            "success": False,
            "error": f"Unsupported action '{action}'. Use 'read' or 'write'.",
        }


def _read_data(path: Path, fmt: str, sheet_name: str | None, encoding: str) -> dict[str, Any]:
    if not path.exists():
        return {
            "success": False,
            "error": f"File does not exist: {path}",
            "file_path": str(path),
        }

    try:
        if fmt == "json":
            with open(path, "r", encoding=encoding, errors="replace") as f:
                parsed = json.load(f)
            return {
                "success": True,
                "format": "json",
                "file_path": str(path),
                "data": parsed,
            }

        elif fmt == "csv":
            with open(path, "r", encoding=encoding, errors="replace", newline="") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
                fields = reader.fieldnames or []
            return {
                "success": True,
                "format": "csv",
                "file_path": str(path),
                "row_count": len(rows),
                "columns": list(fields),
                "data": rows,
            }

        elif fmt == "excel":
            try:
                import openpyxl
            except ImportError:
                return {
                    "success": False,
                    "error": "openpyxl is not installed. Please run 'uv add openpyxl'.",
                    "file_path": str(path),
                }

            wb = openpyxl.load_workbook(str(path), data_only=True)
            target_sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            if target_sheet is None:
                return {"success": False, "error": "No active worksheet found in Excel file."}

            rows_iter = target_sheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if not header_row:
                return {
                    "success": True,
                    "format": "excel",
                    "file_path": str(path),
                    "sheet_name": target_sheet.title,
                    "data": [],
                }

            headers = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(header_row)]
            data_rows = []
            for r in rows_iter:
                row_dict = {headers[i]: r[i] if i < len(r) else None for i in range(len(headers))}
                data_rows.append(row_dict)

            return {
                "success": True,
                "format": "excel",
                "file_path": str(path),
                "sheet_name": target_sheet.title,
                "row_count": len(data_rows),
                "columns": headers,
                "data": data_rows,
            }

        elif fmt == "xml":
            tree = ET.parse(str(path))
            root = tree.getroot()

            def elem_to_dict(elem: ET.Element) -> Any:
                children = list(elem)
                if not children:
                    return elem.text
                d: dict[str, Any] = {}
                for child in children:
                    child_val = elem_to_dict(child)
                    if child.tag in d:
                        if isinstance(d[child.tag], list):
                            d[child.tag].append(child_val)
                        else:
                            d[child.tag] = [d[child.tag], child_val]
                    else:
                        d[child.tag] = child_val
                return d

            return {
                "success": True,
                "format": "xml",
                "file_path": str(path),
                "root_tag": root.tag,
                "data": {root.tag: elem_to_dict(root)},
            }

        return {"success": False, "error": f"Unsupported format: {fmt}"}

    except Exception as exc:
        return {
            "success": False,
            "error": f"Failed reading {fmt} file: {exc}",
            "file_path": str(path),
        }


def _write_data(
    path: Path,
    fmt: str,
    data: Any,
    sheet_name: str | None,
    root_tag: str,
    encoding: str,
) -> dict[str, Any]:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)

        if fmt == "json":
            with open(path, "w", encoding=encoding, errors="replace") as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            return {
                "success": True,
                "action": "write",
                "format": "json",
                "file_path": str(path),
                "bytes_written": path.stat().st_size,
            }

        elif fmt == "csv":
            if not isinstance(data, list):
                return {
                    "success": False,
                    "error": "CSV data must be a list of dictionaries.",
                    "file_path": str(path),
                }

            fields = list(data[0].keys()) if len(data) > 0 and isinstance(data[0], dict) else []
            with open(path, "w", encoding=encoding, newline="") as f:
                writer = csv.DictWriter(f, fieldnames=fields)
                writer.writeheader()
                writer.writerows(data)
            return {
                "success": True,
                "action": "write",
                "format": "csv",
                "file_path": str(path),
                "row_count": len(data),
                "bytes_written": path.stat().st_size,
            }

        elif fmt == "excel":
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            if sheet_name:
                ws.title = sheet_name

            if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
                headers = list(data[0].keys())
                ws.append(headers)
                for item in data:
                    ws.append([item.get(h) for h in headers])
            elif isinstance(data, list):
                for row in data:
                    ws.append(list(row) if isinstance(row, (list, tuple)) else [row])

            wb.save(str(path))
            return {
                "success": True,
                "action": "write",
                "format": "excel",
                "file_path": str(path),
                "bytes_written": path.stat().st_size,
            }

        elif fmt == "xml":
            root = ET.Element(root_tag)

            def dict_to_elem(parent: ET.Element, val: Any) -> None:
                if isinstance(val, dict):
                    for k, v in val.items():
                        child = ET.SubElement(parent, str(k))
                        dict_to_elem(child, v)
                elif isinstance(val, list):
                    for item in val:
                        item_elem = ET.SubElement(parent, "item")
                        dict_to_elem(item_elem, item)
                else:
                    parent.text = str(val) if val is not None else ""

            dict_to_elem(root, data)
            tree = ET.ElementTree(root)
            tree.write(str(path), encoding=encoding, xml_declaration=True)
            return {
                "success": True,
                "action": "write",
                "format": "xml",
                "file_path": str(path),
                "bytes_written": path.stat().st_size,
            }

        return {"success": False, "error": f"Unsupported format: {fmt}"}

    except Exception as exc:
        return {
            "success": False,
            "error": f"Failed writing {fmt} file: {exc}",
            "file_path": str(path),
        }


def execute(**kwargs: Any) -> dict[str, Any]:
    """Standard entry point for dynamic tool execution."""
    return read_write_csv_excel_json_xml(**kwargs)
