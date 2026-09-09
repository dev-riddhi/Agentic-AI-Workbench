"""Utility: Create Word Document (.docx)
Generates styled Word documents from Markdown text using python-docx,
supporting headings, tables, bullet lists, blockquotes, and code snippets,
saving strictly into backend/outputs and registering the file in agent_outputs.
"""

from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def _set_cell_background(cell: Any, hex_color: str) -> None:
    """Sets background color of a docx table cell."""
    clean_hex = hex_color.lstrip("#").upper()
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{clean_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _add_styled_paragraph_with_inline_markdown(p: Any, text: str) -> None:
    """Parses inline bold (**text**) and italic (*text*) into styled runs."""
    # Pattern to capture **bold**, *italic*, and `code`
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(185, 28, 28)
        else:
            p.add_run(token)


def create_docx(
    file_path: str | Path,
    content: str = "",
    title: str | None = None,
    overwrite: bool = True,
    theme_color: str = "#1e40af",
    **kwargs: Any,
) -> dict[str, Any]:
    """Creates a formatted Word (.docx) document inside backend/outputs."""
    backend_dir = Path(__file__).resolve().parents[3]
    outputs_dir = (backend_dir / "outputs").resolve()
    outputs_dir.mkdir(parents=True, exist_ok=True)

    raw_path = str(file_path).strip().replace("\\", "/")
    while raw_path.startswith("./"):
        raw_path = raw_path[2:]
    if raw_path.startswith("outputs/"):
        raw_path = raw_path[len("outputs/") :]
    elif raw_path.startswith("uploads/outputs/"):
        raw_path = raw_path[len("uploads/outputs/") :]
    elif raw_path.startswith("uploads/"):
        raw_path = raw_path[len("uploads/") :]

    p = Path(raw_path)
    dest_path = (outputs_dir / p).resolve()
    if not (dest_path == outputs_dir or dest_path.is_relative_to(outputs_dir)):
        dest_path = (outputs_dir / p.name).resolve()

    if not dest_path.name.lower().endswith(".docx"):
        dest_path = dest_path.with_suffix(".docx")

    if not overwrite and dest_path.exists():
        return {
            "success": False,
            "error": f"File already exists and overwrite is False: {dest_path.name}",
            "file_path": str(dest_path),
        }

    try:
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()

        # Set default page margins: 1 inch
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        lines = content.splitlines()
        i = 0
        total_lines = len(lines)
        detected_title = ""

        clean_hex = theme_color.lstrip("#").upper()
        if len(clean_hex) != 6:
            clean_hex = "1E40AF"
        theme_rgb = RGBColor(int(clean_hex[:2], 16), int(clean_hex[2:4], 16), int(clean_hex[4:], 16))

        while i < total_lines:
            line = lines[i].rstrip()
            stripped = line.strip()

            if not stripped:
                i += 1
                continue

            # Headings
            h1 = re.match(r"^#\s+(.+)$", stripped)
            if h1:
                text = h1.group(1).strip()
                if not detected_title:
                    detected_title = text
                head = doc.add_heading(level=1)
                run = head.add_run(text)
                run.font.color.rgb = theme_rgb
                i += 1
                continue

            h2 = re.match(r"^##\s+(.+)$", stripped)
            if h2:
                text = h2.group(1).strip()
                head = doc.add_heading(level=2)
                run = head.add_run(text)
                run.font.color.rgb = RGBColor(30, 41, 59)
                i += 1
                continue

            h3 = re.match(r"^###\s+(.+)$", stripped)
            if h3:
                text = h3.group(1).strip()
                head = doc.add_heading(level=3)
                head.add_run(text)
                i += 1
                continue

            # Code Blocks
            if stripped.startswith("```"):
                code_lines = []
                i += 1
                while i < total_lines and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                if i < total_lines and lines[i].strip().startswith("```"):
                    i += 1
                p_code = doc.add_paragraph()
                run_code = p_code.add_run("\n".join(code_lines))
                run_code.font.name = "Courier New"
                run_code.font.size = Pt(9)
                run_code.font.color.rgb = RGBColor(15, 23, 42)
                continue

            # Markdown Table
            if "|" in stripped and i + 1 < total_lines and re.match(r"^\|?\s*[\-:]+[\s\-:|]+$", lines[i + 1].strip()):
                header_raw = stripped
                i += 2
                data_rows: list[list[str]] = []
                headers = [c.strip() for c in header_raw.strip("|").split("|")]
                num_cols = len(headers)

                while i < total_lines and "|" in lines[i]:
                    row_raw = lines[i].strip()
                    if not row_raw:
                        break
                    row_cells = [c.strip() for c in row_raw.strip("|").split("|")]
                    while len(row_cells) < num_cols:
                        row_cells.append("")
                    data_rows.append(row_cells[:num_cols])
                    i += 1

                if num_cols > 0:
                    tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                    # Format header row
                    for c_idx, h_text in enumerate(headers):
                        cell = tbl.cell(0, c_idx)
                        _set_cell_background(cell, clean_hex)
                        p_cell = cell.paragraphs[0]
                        run_cell = p_cell.add_run(h_text)
                        run_cell.bold = True
                        run_cell.font.color.rgb = RGBColor(255, 255, 255)
                        run_cell.font.size = Pt(9.5)

                    # Format data rows
                    for r_idx, row_vals in enumerate(data_rows, start=1):
                        for c_idx, cell_val in enumerate(row_vals):
                            cell = tbl.cell(r_idx, c_idx)
                            if r_idx % 2 == 0:
                                _set_cell_background(cell, "F8FAFC")
                            p_cell = cell.paragraphs[0]
                            _add_styled_paragraph_with_inline_markdown(p_cell, cell_val)
                continue

            # Bullet points
            bullet = re.match(r"^(\*|\-|\+)\s+(.+)$", stripped)
            if bullet:
                item_text = bullet.group(2).strip()
                p_bullet = doc.add_paragraph(style="List Bullet")
                _add_styled_paragraph_with_inline_markdown(p_bullet, item_text)
                i += 1
                continue

            # Numbered list
            num = re.match(r"^(\d+)\.\s+(.+)$", stripped)
            if num:
                item_text = num.group(2).strip()
                p_num = doc.add_paragraph(style="List Number")
                _add_styled_paragraph_with_inline_markdown(p_num, item_text)
                i += 1
                continue

            # Blockquote
            if stripped.startswith(">"):
                quote_text = stripped.lstrip("> ").strip()
                p_quote = doc.add_paragraph()
                p_quote.paragraph_format.left_indent = Inches(0.5)
                run_quote = p_quote.add_run(f"“ {quote_text} ”")
                run_quote.italic = True
                run_quote.font.color.rgb = RGBColor(71, 85, 105)
                i += 1
                continue

            # Standard body paragraph
            p_body = doc.add_paragraph()
            _add_styled_paragraph_with_inline_markdown(p_body, stripped)
            i += 1

        doc.save(str(dest_path))
        file_size = dest_path.stat().st_size
        doc_title = title or detected_title or dest_path.stem.replace("_", " ").title()

        # Update agent_outputs db record
        try:
            from app.runtime.agent_runtime import record_tool_generated_output
            record_tool_generated_output(
                file_path=str(dest_path),
                title=doc_title,
                output_type="docx",
                content_preview=content[:500] if content else None,
            )
        except Exception:
            pass

        return {
            "success": True,
            "file_path": str(dest_path),
            "filename": dest_path.name,
            "file_size": file_size,
            "title": doc_title,
            "output_type": "docx",
            "message": f"Successfully created Word document '{dest_path.name}' ({file_size} bytes) in outputs directory.",
        }
    except Exception as exc:
        return {
            "success": False,
            "error": f"Failed to generate Word document: {exc}",
            "file_path": str(dest_path),
            "output_type": "docx",
        }
