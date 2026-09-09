"""Utility: Create Excel (.xlsx)
Generates styled Excel workbooks from tabular data, Markdown tables, or CSV content using openpyxl,
saving strictly into backend/outputs and registering the file in agent_outputs.
"""

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.tools.utils.create_csv import parse_content_to_rows


def create_excel(
    file_path: str | Path,
    content: str = "",
    title: str | None = None,
    sheet_name: str = "Sheet1",
    overwrite: bool = True,
    theme_color: str = "1E40AF",
    **kwargs: Any,
) -> dict[str, Any]:
    """Creates a beautifully styled Excel (.xlsx) spreadsheet inside backend/outputs."""
    backend_dir = Path(__file__).resolve().parents[3]
    outputs_dir = (backend_dir / "outputs").resolve()
    outputs_dir.mkdir(parents=True, exist_ok=True)

    raw_path = str(file_path).strip().replace("\\", "/")
    while raw_path.startswith("./"):
        raw_path = raw_path[2:]
    if raw_path.startswith("outputs/"):
        raw_path = raw_path[len("outputs/") :]
    elif raw_path.startswith("uploads/outputs/"):
        raw_path = raw_path[len("uploads/outputs/") :]
    elif raw_path.startswith("uploads/"):
        raw_path = raw_path[len("uploads/") :]

    p = Path(raw_path)
    dest_path = (outputs_dir / p).resolve()
    if not (dest_path == outputs_dir or dest_path.is_relative_to(outputs_dir)):
        dest_path = (outputs_dir / p.name).resolve()

    if not dest_path.name.lower().endswith(".xlsx"):
        dest_path = dest_path.with_suffix(".xlsx")

    if not overwrite and dest_path.exists():
        return {
            "success": False,
            "error": f"File already exists and overwrite is False: {dest_path.name}",
            "file_path": str(dest_path),
        }

    try:
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        rows = parse_content_to_rows(content)

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31] if sheet_name else "Sheet1"

        # Theme styling definitions
        clean_theme = theme_color.lstrip("#").upper()
        if len(clean_theme) != 6:
            clean_theme = "1E40AF"

        header_fill = PatternFill(start_color=clean_theme, end_color=clean_theme, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        thin_side = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # Write data rows
        if rows:
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, val in enumerate(row, start=1):
                    # Try numeric conversion if applicable
                    cell_val: Any = val
                    if isinstance(val, str):
                        clean_num = val.replace(",", "").strip()
                        if clean_num.isdigit():
                            cell_val = int(clean_num)
                        else:
                            try:
                                cell_val = float(clean_num)
                            except ValueError:
                                cell_val = val

                    cell = ws.cell(row=r_idx, column=c_idx, value=cell_val)
                    cell.border = cell_border

                    if r_idx == 1:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    else:
                        cell.font = regular_font
                        cell.alignment = Alignment(vertical="center")
                        if r_idx % 2 == 0:
                            cell.fill = alt_fill
        else:
            ws.cell(row=1, column=1, value=content or "No tabular data provided.")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(str(dest_path))
        file_size = dest_path.stat().st_size
        doc_title = title or dest_path.stem.replace("_", " ").title()

        # Update agent_outputs db record
        try:
            from app.runtime.agent_runtime import record_tool_generated_output
            record_tool_generated_output(
                file_path=str(dest_path),
                title=doc_title,
                output_type="xlsx",
                content_preview=content[:500] if content else None,
            )
        except Exception:
            pass

        return {
            "success": True,
            "file_path": str(dest_path),
            "filename": dest_path.name,
            "file_size": file_size,
            "title": doc_title,
            "output_type": "xlsx",
            "rows_count": len(rows),
            "message": f"Successfully created Excel workbook '{dest_path.name}' ({len(rows)} rows, {file_size} bytes) in outputs directory.",
        }
    except Exception as exc:
        return {
            "success": False,
            "error": f"Failed to generate Excel workbook: {exc}",
            "file_path": str(dest_path),
            "output_type": "xlsx",
        }
